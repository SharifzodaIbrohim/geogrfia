"""Admin XLSX export — /api/admin/export/* (openpyxl). Safe route rebind."""
from __future__ import annotations

import io
import logging
from datetime import datetime, timezone
from typing import Any

log = logging.getLogger("geografia.export")

STATUS_TJ = {
    "passed": "Гузашт",
    "pass": "Гузашт",
    "failed": "Нагузашт",
    "fail": "Нагузашт",
    "timeout": "Вақт тамом",
    "submitted": "Супорида шуд",
    "in_progress": "Дар ҷараён",
}


def _status_label(st: Any) -> str:
    s = str(st or "").lower()
    return STATUS_TJ.get(s, str(st or "—"))


def _sid(row: dict) -> str:
    v = (
        row.get("studentCode")
        or row.get("studentId")
        or row.get("student_id")
        or row.get("id")
        or ""
    )
    return str(v)


def _name(row: dict) -> str:
    return str(
        row.get("fullName")
        or row.get("studentName")
        or row.get("name")
        or row.get("full_name")
        or "—"
    )


def _school(row: dict) -> str:
    return str(row.get("school") or row.get("studentSchool") or row.get("schoolName") or "—")


def _class(row: dict) -> str:
    return str(row.get("className") or row.get("studentClass") or row.get("class_name") or "—")


def _oly(row: dict) -> str:
    return str(row.get("olympiadTitle") or row.get("title") or row.get("olympiad_title") or "—")


def _score_pct(row: dict) -> Any:
    if row.get("score") is not None:
        return row.get("score")
    c, t = row.get("correct"), row.get("total")
    try:
        if t and float(t) > 0:
            return round(100.0 * float(c or 0) / float(t), 1)
    except Exception:
        pass
    return ""


def _style_header(ws, headers):
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    fill = PatternFill("solid", fgColor="1B4332")
    font = Font(bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin
    return thin


def _wb_results(rows: list) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Натиҷаҳо"
    headers = [
        "Ном",
        "Student ID",
        "Мактаб",
        "Синф",
        "Олимпиада",
        "Хол (%)",
        "Дуруст",
        "Ҳама",
        "Статус",
        "Санаи анҷом",
    ]
    thin = _style_header(ws, headers)
    for r_i, row in enumerate(rows, 2):
        vals = [
            _name(row),
            _sid(row),
            _school(row),
            _class(row),
            _oly(row),
            _score_pct(row),
            row.get("correct") if row.get("correct") is not None else "",
            row.get("total") if row.get("total") is not None else "",
            _status_label(row.get("status")),
            str(row.get("finishedAt") or row.get("finished_at") or "")[:19],
        ]
        for c_i, v in enumerate(vals, 1):
            cell = ws.cell(r_i, c_i, v)
            cell.border = thin
            if c_i == 2:
                cell.number_format = "@"
                cell.alignment = Alignment(horizontal="left")
    for i, w in enumerate([28, 22, 18, 10, 24, 10, 10, 10, 14, 20], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _wb_students(rows: list) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Хонандагон"
    headers = ["Student ID", "Ном", "Мактаб", "Синф", "Санаи бақайд"]
    thin = _style_header(ws, headers)
    for r_i, row in enumerate(rows, 2):
        vals = [
            str(row.get("id") or row.get("studentId") or ""),
            str(row.get("fullName") or row.get("name") or "—"),
            str(row.get("school") or "—"),
            str(row.get("className") or "—"),
            str(row.get("createdAt") or "")[:19],
        ]
        for c_i, v in enumerate(vals, 1):
            cell = ws.cell(r_i, c_i, v)
            cell.border = thin
            if c_i == 1:
                cell.number_format = "@"
    for i, w in enumerate([22, 28, 18, 10, 20], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _wb_full(results: list, students: list, olympiads: list) -> io.BytesIO:
    """Single workbook, 3 sheets — no private openpyxl APIs."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    ws = wb.active
    ws.title = "Натиҷаҳо"
    headers = [
        "Ном",
        "Student ID",
        "Мактаб",
        "Синф",
        "Олимпиада",
        "Хол (%)",
        "Дуруст",
        "Ҳама",
        "Статус",
        "Санаи анҷом",
    ]
    thin = _style_header(ws, headers)
    for r_i, row in enumerate(results, 2):
        vals = [
            _name(row),
            _sid(row),
            _school(row),
            _class(row),
            _oly(row),
            _score_pct(row),
            row.get("correct") if row.get("correct") is not None else "",
            row.get("total") if row.get("total") is not None else "",
            _status_label(row.get("status")),
            str(row.get("finishedAt") or row.get("finished_at") or "")[:19],
        ]
        for c_i, v in enumerate(vals, 1):
            cell = ws.cell(r_i, c_i, v)
            cell.border = thin
            if c_i == 2:
                cell.number_format = "@"
                cell.alignment = Alignment(horizontal="left")
    for i, w in enumerate([28, 22, 18, 10, 24, 10, 10, 10, 14, 20], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Хонандагон")
    h2 = ["Student ID", "Ном", "Мактаб", "Синф", "Санаи бақайд"]
    thin2 = _style_header(ws2, h2)
    for r_i, row in enumerate(students, 2):
        vals = [
            str(row.get("id") or row.get("studentId") or ""),
            str(row.get("fullName") or row.get("name") or "—"),
            str(row.get("school") or "—"),
            str(row.get("className") or "—"),
            str(row.get("createdAt") or "")[:19],
        ]
        for c_i, v in enumerate(vals, 1):
            cell = ws2.cell(r_i, c_i, v)
            cell.border = thin2
            if c_i == 1:
                cell.number_format = "@"
    for i, w in enumerate([22, 28, 18, 10, 20], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("Олимпиадаҳо")
    h3 = ["ID", "Унвон", "Навъ", "Ҳад %", "Фаъол", "Сана"]
    thin3 = _style_header(ws3, h3)
    for r_i, o in enumerate(olympiads, 2):
        vals = [
            str(o.get("id") or ""),
            o.get("title") or "",
            o.get("type") or "",
            o.get("passScore") or o.get("pass_score") or "",
            "Ҳа" if o.get("active") else "Не",
            str(o.get("createdAt") or "")[:19],
        ]
        for c_i, v in enumerate(vals, 1):
            cell = ws3.cell(r_i, c_i, v)
            cell.border = thin3
            if c_i == 1:
                cell.number_format = "@"
    for i, w in enumerate([36, 28, 12, 10, 10, 20], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _filter_rows(rows: list, args) -> list:
    school = (args.get("school") or "").strip().lower()
    class_name = (args.get("className") or "").strip().lower()
    status = (args.get("status") or "").strip().lower()
    q = (args.get("q") or "").strip().lower()
    oid = (args.get("olympiadId") or "").strip()
    try:
        smin = float(args.get("scoreMin")) if args.get("scoreMin") not in (None, "") else None
    except Exception:
        smin = None
    try:
        smax = float(args.get("scoreMax")) if args.get("scoreMax") not in (None, "") else None
    except Exception:
        smax = None

    out = []
    for r in rows:
        if oid and str(r.get("olympiadId") or "") != oid:
            continue
        if school and school not in _school(r).lower():
            continue
        if class_name and class_name not in _class(r).lower():
            continue
        if status and status not in str(r.get("status") or "").lower():
            continue
        if q:
            blob = " ".join([_name(r), _sid(r), _school(r), _class(r), _oly(r)]).lower()
            if q not in blob:
                continue
        sc = _score_pct(r)
        try:
            scf = float(sc) if sc != "" and sc is not None else None
        except Exception:
            scf = None
        if smin is not None and (scf is None or scf < smin):
            continue
        if smax is not None and (scf is None or scf > smax):
            continue
        out.append(r)
    return out


def install(app=None):
    from flask import request, send_file, jsonify

    if app is None:
        from flask import current_app

        app = current_app._get_current_object()

    def _auth_headers():
        h = {}
        tok = request.headers.get("X-Admin-Token") or ""
        auth = request.headers.get("Authorization") or ""
        if tok:
            h["X-Admin-Token"] = tok
        if auth:
            h["Authorization"] = auth
        return h

    def _get_json(path: str):
        with app.test_client() as client:
            for k, v in request.cookies.items():
                try:
                    client.set_cookie("localhost", k, v)
                except Exception:
                    pass
            r = client.get(path, headers=_auth_headers())
            if r.status_code == 401:
                return None, 401
            if r.status_code >= 400:
                log.warning("export internal GET %s -> %s", path, r.status_code)
                return None, r.status_code
            try:
                return r.get_json(silent=True) or {}, 200
            except Exception:
                return {}, 200

    def _send_xlsx(buf: io.BytesIO, filename: str):
        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename,
        )

    def admin_export_preview():
        try:
            data, code = _get_json("/api/admin/results?" + request.query_string.decode())
            if code == 401:
                return jsonify({"error": "Дастрасӣ рад шуд."}), 401
            rows = (data or {}).get("results") or (data or {}).get("items") or []
            rows = _filter_rows(rows, request.args)
            return jsonify({"count": len(rows), "total": len(rows)})
        except Exception as e:
            log.exception("export preview")
            return jsonify({"error": str(e)}), 500

    def admin_export_results():
        try:
            data, code = _get_json("/api/admin/results?" + request.query_string.decode())
            if code == 401:
                return jsonify({"error": "Дастрасӣ рад шуд."}), 401
            rows = (data or {}).get("results") or (data or {}).get("items") or []
            rows = _filter_rows(rows, request.args)
            buf = _wb_results(rows)
            stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M")
            return _send_xlsx(buf, f"Geografia_Results_{stamp}.xlsx")
        except Exception as e:
            log.exception("export results")
            return jsonify({"error": "Export хато: " + str(e)}), 500

    def admin_export_olympiad(oid):
        try:
            data, code = _get_json(f"/api/admin/olympiads/{oid}/results")
            if code == 401:
                return jsonify({"error": "Дастрасӣ рад шуд."}), 401
            if code >= 400 or not data:
                data, code = _get_json(f"/api/admin/results?olympiadId={oid}")
                if code == 401:
                    return jsonify({"error": "Дастрасӣ рад шуд."}), 401
            rows = (data or {}).get("results") or (data or {}).get("items") or []
            buf = _wb_results(rows)
            stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M")
            return _send_xlsx(buf, f"Geografia_Olympiad_{stamp}.xlsx")
        except Exception as e:
            log.exception("export olympiad")
            return jsonify({"error": "Export хато: " + str(e)}), 500

    def admin_export_students():
        try:
            data, code = _get_json("/api/admin/students")
            if code == 401:
                return jsonify({"error": "Дастрасӣ рад шуд."}), 401
            rows = (data or {}).get("students") or (data or {}).get("items") or []
            buf = _wb_students(rows)
            stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M")
            return _send_xlsx(buf, f"Geografia_Students_{stamp}.xlsx")
        except Exception as e:
            log.exception("export students")
            return jsonify({"error": "Export хато: " + str(e)}), 500

    def admin_export_full():
        try:
            res_data, code = _get_json("/api/admin/results")
            if code == 401:
                return jsonify({"error": "Дастрасӣ рад шуд."}), 401
            st_data, _ = _get_json("/api/admin/students")
            ol_data, _ = _get_json("/api/admin/olympiads")
            results = (res_data or {}).get("results") or []
            students = (st_data or {}).get("students") or []
            olympiads = (ol_data or {}).get("olympiads") or (
                ol_data if isinstance(ol_data, list) else []
            )
            buf = _wb_full(results, students, olympiads)
            stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M")
            return _send_xlsx(buf, f"Geografia_Full_{stamp}.xlsx")
        except Exception as e:
            log.exception("export full")
            return jsonify({"error": "Export хато: " + str(e)}), 500

    routes = [
        ("/api/admin/export/preview", "geo_export_preview", admin_export_preview, ["GET"]),
        ("/api/admin/export/results", "geo_export_results", admin_export_results, ["GET"]),
        ("/api/admin/export/olympiad/<oid>", "geo_export_olympiad", admin_export_olympiad, ["GET"]),
        ("/api/admin/export/students", "geo_export_students", admin_export_students, ["GET"]),
        ("/api/admin/export/full", "geo_export_full", admin_export_full, ["GET"]),
    ]

    for rule, endpoint, view, methods in routes:
        if endpoint in app.view_functions:
            try:
                del app.view_functions[endpoint]
            except Exception:
                pass
        try:
            app.add_url_rule(rule, endpoint=endpoint, view_func=view, methods=methods)
        except Exception as e:
            log.warning("add_url_rule %s failed: %s", rule, e)
            app.view_functions[endpoint] = view
            try:
                app.add_url_rule(
                    rule, endpoint=endpoint, view_func=view, methods=methods, strict_slashes=False
                )
            except Exception as e2:
                log.error("export route bind failed %s: %s", rule, e2)

    for ep, fn in (
        ("admin_export_students", admin_export_students),
        ("export_students", admin_export_students),
        ("admin_export_full", admin_export_full),
    ):
        if ep in app.view_functions:
            app.view_functions[ep] = fn

    print("[boot] patch_admin_export v2: XLSX /api/admin/export/* bound")
    return True
